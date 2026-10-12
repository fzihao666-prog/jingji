#!/usr/bin/env python3
"""仅清理显式文件；不输出单元格、字段值或异常详情。"""
import argparse
import csv
import io
import json
import os
from pathlib import Path
import re
import shutil
import tempfile
import zipfile
from copy import copy

KEYS = {'身份证', '身份证号', '身份证号码', '证件号码', 'idcard', 'idcardnumber', 'identitynumber', 'identitycard', 'identitycardnumber'}
def identity(value):
    return isinstance(value, str) and re.sub(r'[\s_\-]', '', value).lower() in KEYS

def clean_json(value):
    count = 0
    if isinstance(value, dict):
        result = {}
        for key, item in value.items():
            if identity(key):
                count += 1
            else:
                result[key], hits = clean_json(item)
                count += hits
        return result, count
    if isinstance(value, list):
        # 二维表格只按首行的明确字段名删除整列。
        if value and isinstance(value[0], list):
            columns = [i for i, item in enumerate(value[0]) if identity(item)]
            if columns:
                result, count = [], len(columns)
                for row in value:
                    if isinstance(row, list):
                        cleaned_row = []
                        for index, item in enumerate(row):
                            if index not in columns:
                                cleaned_item, hits = clean_json(item)
                                cleaned_row.append(cleaned_item)
                                count += hits
                        result.append(cleaned_row)
                    else:
                        cleaned_row, hits = clean_json(row)
                        result.append(cleaned_row)
                        count += hits
                return result, count
        result = []
        for item in value:
            item, hits = clean_json(item)
            result.append(item)
            count += hits
        return result, count
    return value, 0

def shifted_range(reference, column):
    from openpyxl.worksheet.cell_range import CellRange
    target = CellRange(str(reference))
    if target.min_col == target.max_col == column:
        return None
    target.min_col -= int(target.min_col > column)
    target.max_col -= int(target.max_col >= column)
    return str(target)


def safe_validation(validation):
    if validation.type == 'list':
        return isinstance(validation.formula1, str) and re.fullmatch(r'"[^"\r\n]*"', validation.formula1) is not None and validation.formula2 is None
    if validation.type == 'date':
        return all(value is None or re.fullmatch(r'[0-9]+(?:\.[0-9]+)?', str(value)) is not None for value in (validation.formula1, validation.formula2))
    return False


def transform(path):
    suffix = path.suffix.lower()
    if suffix == '.xlsx':
        from openpyxl import load_workbook
        # openpyxl无法可靠维护这些对象引用；预览和执行均拒绝。
        with zipfile.ZipFile(path) as archive:
            names = archive.namelist()
            forbidden = ('xl/drawings/', 'xl/charts/', 'xl/tables/', 'xl/pivot', 'xl/externalLinks/', 'xl/activeX/', 'customXml/', '_xmlsignatures/', 'xl/slicers/', 'xl/ctrlProps/', 'xl/embeddings/')
            if any(name.startswith(forbidden) or 'vbaProject' in name for name in names):
                raise ValueError('复杂工作簿对象不支持安全列删除')
            if any(b'extLst' in archive.read(name) for name in names if name.startswith('xl/') and name.endswith('.xml')):
                raise ValueError('复杂工作簿对象不支持安全列删除')
        workbook = load_workbook(path)
        count = 0
        for sheet in workbook:
            if sheet.tables or sheet.conditional_formatting or sheet.auto_filter.ref or any(not safe_validation(rule) for rule in sheet.data_validations.dataValidation):
                raise ValueError('复杂工作簿对象不支持安全列删除')
            if any(cell.data_type == 'f' for row in sheet for cell in row) or workbook.defined_names:
                raise ValueError('复杂工作簿对象不支持安全列删除')
            columns = {cell.column for row in sheet.iter_rows(min_row=1, max_row=min(sheet.max_row, 20)) for cell in row if identity(cell.value)}
            # 说明单元格不作为整列依据，仅删除明确提及身份证的说明内容。
            for row in sheet:
                for cell in row:
                    if isinstance(cell.value, str) and '身份证' in cell.value and not identity(cell.value):
                        if cell.value.strip() == '身份证为18位，末位可为X':
                            cell.value = '生日和性别请独立填写'
                            count += 1
                        else:
                            raise ValueError('混合文本含身份证说明需人工处理')
            for column in sorted(columns, reverse=True):
                from openpyxl.utils import get_column_letter, column_index_from_string
                dimensions = [(column_index_from_string(key), value) for key, value in sheet.column_dimensions.items()]
                # 合并范围和验证范围不会被delete_cols维护，显式转换引用。
                merged = list(sheet.merged_cells.ranges)
                styles = {(cell.row, cell.column): copy(cell._style) for row in sheet for cell in row}
                for region in merged:
                    if region.min_col == column and region.max_col > column and sheet.cell(region.min_row, column).value is not None:
                        raise ValueError('复杂工作簿对象不支持安全列删除')
                    sheet.unmerge_cells(str(region))
                sheet.delete_cols(column)
                for region in merged:
                    reference = shifted_range(region, column)
                    if reference:
                        sheet.merge_cells(reference)
                for (row, index), style in styles.items():
                    if index != column:
                        sheet.cell(row, index - int(index > column))._style = style
                for rule in list(sheet.data_validations.dataValidation):
                    references = [shifted_range(region, column) for region in rule.sqref.ranges]
                    references = [reference for reference in references if reference]
                    if references:
                        rule.sqref = ' '.join(references)
                    else:
                        sheet.data_validations.dataValidation.remove(rule)
                sheet.column_dimensions.clear()
                for index, dimension in dimensions:
                    if index != column:
                        key = get_column_letter(index - (index > column))
                        dimension.index = key
                        dimension.min = dimension.max = index - (index > column)
                        sheet.column_dimensions[key] = dimension
                count += 1
        output = io.BytesIO()
        workbook.save(output)
        return output.getvalue(), count
    if suffix not in {'.csv', '.json', '.ndjson', '.jsonl', '.log', '.txt', '.xml'}:
        raise ValueError('不支持的文件类型')
    text = path.read_text(encoding='utf-8-sig')
    if suffix == '.csv':
        rows = list(csv.reader(io.StringIO(text)))
        columns = [i for i, item in enumerate(rows[0] if rows else []) if identity(item)]
        output = io.StringIO(newline='')
        csv.writer(output).writerows([[item for i, item in enumerate(row) if i not in columns] for row in rows])
        return output.getvalue().encode(), len(columns)
    if suffix in {'.json', '.ndjson', '.jsonl'}:
        if suffix == '.json':
            value, count = clean_json(json.loads(text))
            return (json.dumps(value, ensure_ascii=False, indent=2) + '\n').encode(), count
        output, count = [], 0
        for line in text.splitlines():
            if line.strip():
                value, hits = clean_json(json.loads(line))
                output.append(json.dumps(value, ensure_ascii=False))
                count += hits
        return ('\n'.join(output) + '\n').encode(), count
    if suffix == '.xml':
        raise ValueError('XML结构文件需人工确认节点语义，拒绝自动修改')
    output, count = [], 0
    for line in text.splitlines(keepends=True):
        # 混合日志不能整行删除业务字段；JSON日志按字段清理，否则只删独立键值行。
        try:
            value, hits = clean_json(json.loads(line))
            output.append(json.dumps(value, ensure_ascii=False) + '\n')
            count += hits
        except (ValueError, TypeError):
            match = re.match(r'^\s*([^:=：]+)\s*[:=：]', line)
            if match and identity(match.group(1)):
                # 仅支持独立标量；分隔符或其它字段语义无法证明时拒绝。
                if re.fullmatch(r'[A-Za-z0-9_-]*', line[match.end():].strip()) is None:
                    raise ValueError('混合文本含身份证说明需人工处理')
                count += 1
            elif '身份证' in line:
                raise ValueError('混合文本含身份证说明需人工处理')
            else:
                output.append(line)
    return ''.join(output).encode(), count

def apply(path, content, backup_dir):
    backup_dir.mkdir(mode=0o700, parents=True, exist_ok=True)
    if backup_dir.is_symlink() or backup_dir.stat().st_mode & 0o077:
        raise ValueError('备份目录权限必须为0700')
    backup = backup_dir / path.name
    backup_descriptor = os.open(backup, os.O_WRONLY | os.O_CREAT | os.O_EXCL, 0o600)
    with os.fdopen(backup_descriptor, 'wb') as handle:
        os.fchmod(handle.fileno(), 0o600)
        with path.open('rb') as source:
            shutil.copyfileobj(source, handle)
        handle.flush()
        os.fsync(handle.fileno())
    descriptor, temporary = tempfile.mkstemp(prefix='.identity-clean-', dir=path.parent)
    try:
        with os.fdopen(descriptor, 'wb') as handle:
            os.fchmod(handle.fileno(), path.stat().st_mode & 0o777)
            handle.write(content)
            handle.flush()
            os.fsync(handle.fileno())
        os.replace(temporary, path)
    finally:
        if os.path.exists(temporary):
            os.unlink(temporary)

def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--file', action='append', required=True)
    parser.add_argument('--apply', action='store_true')
    parser.add_argument('--confirmed', action='store_true')
    parser.add_argument('--backup-dir', type=Path)
    args = parser.parse_args()
    if args.apply and (not args.confirmed or not args.backup_dir):
        parser.error('--apply要求--confirmed和--backup-dir')
    failed = False
    for name in args.file:
        path = Path(name)
        try:
            if path.is_symlink() or not path.is_file():
                raise ValueError('仅支持普通文件')
            content, count = transform(path)
            if args.apply and count:
                apply(path, content, args.backup_dir)
            print(json.dumps({'path': str(path), 'hits': count, 'applied': bool(args.apply and count)}, ensure_ascii=False))
        except Exception as error:
            failed = True
            reason = str(error) if isinstance(error, ValueError) and str(error) in {'复杂工作簿对象不支持安全列删除', '不支持的文件类型', 'XML结构文件需人工确认节点语义，拒绝自动修改', '混合文本含身份证说明需人工处理', '备份目录权限必须为0700', '仅支持普通文件'} else '文件格式或安全写入检查失败；源文件未替换'
            print(json.dumps({'path': str(path), 'error': reason}, ensure_ascii=False))
    return int(failed)
if __name__ == '__main__':
    raise SystemExit(main())
